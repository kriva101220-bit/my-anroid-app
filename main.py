import sqlite3
import tkinter as tk
from tkinter import ttk, messagebox
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# --- ૧. ડેટાબેઝ સેટઅપ ---
conn = sqlite3.connect("customers.db")
cursor = conn.cursor()

cursor.execute("""
CREATE TABLE IF NOT EXISTS customers (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    name TEXT,
    phone TEXT,
    home_no TEXT,
    amount REAL
)
""")
conn.commit()

# જૂના ડેટાબેઝમાં નવી કોલમ ઓટોમેટિક ઉમેરવા માટે
try:
    cursor.execute("ALTER TABLE customers ADD COLUMN home_no TEXT")
    conn.commit()
except sqlite3.OperationalError:
    pass

try:
    cursor.execute("ALTER TABLE customers ADD COLUMN amount REAL")
    conn.commit()
except sqlite3.OperationalError:
    pass

# --- ૨. ફંકશન્સ ---

# કસ્ટમર ઉમેરવો
def add_customer():
    name = entry_name.get()
    phone = entry_phone.get()
    home_no = entry_home.get()
    amount = entry_amount.get()

    if name == "" or phone == "":
        messagebox.showwarning("ચેતવણી", "કૃપા કરીને નામ અને મોબાઇલ નંબર લખો!")
        return

    if amount == "":
        amount = 0.0
    else:
        try:
            amount = float(amount)
        except ValueError:
            messagebox.showwarning("ચેતવણી", "રકમ ફક્ત નંબરમાં જ લખો!")
            return

    cursor.execute("INSERT INTO customers (name, phone, home_no, amount) VALUES (?, ?, ?, ?)", (name, phone, home_no, amount))
    conn.commit()
    messagebox.showinfo("સફળતા", "કસ્ટમરની એન્ટ્રી સેવ થઈ ગઈ છે!")
    
    entry_name.delete(0, tk.END)
    entry_phone.delete(0, tk.END)
    entry_home.delete(0, tk.END)
    entry_amount.delete(0, tk.END)
    
    show_customers()

# ડેટા બતાવવો અને સરવાળો ગણવો
def show_customers(rows=None):
    for row in tree.get_children():
        tree.delete(row)
    
    if rows is None:
        cursor.execute("SELECT id, name, phone, home_no, amount FROM customers")
        rows = cursor.fetchall()
        
    total_sum = 0.0

    for index, row in enumerate(rows, start=1):
        db_id = row[0]
        amt = row[4] if row[4] is not None else 0.0
        total_sum += float(amt)
        tree.insert("", tk.END, iid=str(db_id), values=(index, row[1], row[2], row[3], amt))

    label_total.config(text=f"કુલ રકમ: ₹ {total_sum:,.2f}")

# ડેટા ડિલીટ કરવો
def delete_customer():
    selected_item = tree.selection()
    if not selected_item:
        messagebox.showwarning("ચેતવણી", "ડિલીટ કરવા માટે પહેલાં લિસ્ટમાંથી કસ્ટમર પસંદ કરો!")
        return
    
    confirm = messagebox.askyesno("ખરાઈ કરો", "શું તમે ખરેખર આ કસ્ટમરને ડિલીટ કરવા માંગો છો?")
    if confirm:
        customer_id = selected_item[0]
        
        cursor.execute("DELETE FROM customers WHERE id=?", (customer_id,))
        conn.commit()
        messagebox.showinfo("સફળતા", "કસ્ટમર ડિલીટ થઈ ગયા છે!")
        show_customers()

# ડેટા શોધો (Search)
def search_customer():
    query = entry_search.get()
    if query == "":
        show_customers()
    else:
        cursor.execute("SELECT id, name, phone, home_no, amount FROM customers WHERE name LIKE ? OR phone LIKE ? OR home_no LIKE ?", ('%' + query + '%', '%' + query + '%', '%' + query + '%'))
        rows = cursor.fetchall()
        show_customers(rows)

# ડિઝાઇનવાળી પ્રોફેસનલ એક્સેલ ફાઇલ એક્સપોર્ટ કરવી
def export_to_excel():
    cursor.execute("SELECT id, name, phone, home_no, amount FROM customers")
    rows = cursor.fetchall()
    
    if not rows:
        messagebox.showwarning("ચેતવણી", "એક્સપોર્ટ કરવા માટે કોઈ ડેટા નથી!")
        return

    # નવી એક્સેલ વર્કબુક બનાવો
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Customer Report"

    # સ્ટાઇલ ઓબ્જેક્ટ્સ
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")  # ઘાટો વાદળી રંગ
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    data_font = Font(name="Calibri", size=11)
    total_font = Font(name="Calibri", size=11, bold=True)
    total_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")  # આછો વાદળી રંગ

    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )

    align_center = Alignment(horizontal='center', vertical='center')
    align_left = Alignment(horizontal='left', vertical='center')
    align_right = Alignment(horizontal='right', vertical='center')

    # હેડર ઉમેરો
    headers = ["ક્રમ નં.", "કસ્ટમરનું નામ", "મોબાઇલ નંબર", "ઘર નંબર (Home No)", "રકમ (RS)"]
    ws.append(headers)

    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = align_center
        cell.border = thin_border

    # ડેટા ઉમેરો
    start_row = 2
    for index, row in enumerate(rows, start=1):
        amt = row[4] if row[4] is not None else 0.0
        ws.append([index, row[1], row[2], row[3], amt])
        
        current_row = start_row + index - 1
        ws.cell(row=current_row, column=1).alignment = align_center
        ws.cell(row=current_row, column=2).alignment = align_left
        ws.cell(row=current_row, column=3).alignment = align_center
        ws.cell(row=current_row, column=4).alignment = align_center
        
        amt_cell = ws.cell(row=current_row, column=5)
        amt_cell.alignment = align_right
        amt_cell.number_format = '#,##0.00'

        for col_num in range(1, len(headers) + 1):
            c = ws.cell(row=current_row, column=col_num)
            c.font = data_font
            c.border = thin_border

    # ટોટલ રો (સરવાળાની લાઇન) ઉમેરો
    last_data_row = len(rows) + 1
    total_row = last_data_row + 1

    ws.cell(row=total_row, column=4, value="કુલ રકમ (Total):").alignment = align_right
    ws.cell(row=total_row, column=4).font = total_font

    total_cell = ws.cell(row=total_row, column=5, value=f"=SUM(E2:E{last_data_row})")
    total_cell.alignment = align_right
    total_cell.font = total_font
    total_cell.number_format = '₹ #,##0.00'

    for col_num in range(1, len(headers) + 1):
        c = ws.cell(row=total_row, column=col_num)
        c.fill = total_fill
        c.border = thin_border

    # ઓટોમેટિક કોલમ પહોળાઈ ગોઠવવી
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value or '')
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = max(max_len + 5, 12)

    filename = "Customer_Data_Report.xlsx"
    try:
        wb.save(filename)
        messagebox.showinfo("સફળતા", f"ડિઝાઇનવાળી એક્સેલ ફાઇલ '{filename}' તરીકે સેવ થઈ ગઈ છે!")
    except Exception as e:
        messagebox.showerror("ભૂલ", f"ફાઈલ સેવ કરવામાં ભૂલ આવી: {e}")

# --- ૩. સ્ક્રીન (GUI) ડિઝાઇન ---
root = tk.Tk()
root.title("કસ્ટમર મેનેજમેન્ટ સોફ્ટવેર")
root.state('zoomed')

# ટાઇટલ
title_label = tk.Label(root, text="કસ્ટમર એન્ટ્રી અને મેનેજમેન્ટ સોફ્ટવેર", font=("Arial", 16, "bold"), bg="#1a73e8", fg="white", pady=10)
title_label.pack(fill=tk.X)

# ઇનપુટ ફોર્મ ફ્રેમ
frame_form = tk.LabelFrame(root, text=" નવી એન્ટ્રી કરો ", font=("Arial", 12, "bold"), padx=15, pady=15)
frame_form.pack(pady=15, fill=tk.X, padx=20)

tk.Label(frame_form, text="કસ્ટમરનું નામ:", font=("Arial", 11)).grid(row=0, column=0, padx=10, pady=5, sticky="w")
entry_name = tk.Entry(frame_form, font=("Arial", 11), width=20)
entry_name.grid(row=0, column=1, padx=10, pady=5)

tk.Label(frame_form, text="મોબાઇલ નંબર:", font=("Arial", 11)).grid(row=0, column=2, padx=10, pady=5, sticky="w")
entry_phone = tk.Entry(frame_form, font=("Arial", 11), width=18)
entry_phone.grid(row=0, column=3, padx=10, pady=5)

tk.Label(frame_form, text="ઘર નંબર (Home No):", font=("Arial", 11)).grid(row=0, column=4, padx=10, pady=5, sticky="w")
entry_home = tk.Entry(frame_form, font=("Arial", 11), width=15)
entry_home.grid(row=0, column=5, padx=10, pady=5)

tk.Label(frame_form, text="રકમ (RS):", font=("Arial", 11)).grid(row=0, column=6, padx=10, pady=5, sticky="w")
entry_amount = tk.Entry(frame_form, font=("Arial", 11), width=15)
entry_amount.grid(row=0, column=7, padx=10, pady=5)

btn_add = tk.Button(frame_form, text="સેવ કરો", command=add_customer, bg="#28a745", fg="white", font=("Arial", 11, "bold"), width=12)
btn_add.grid(row=0, column=8, padx=15, pady=5)

# સર્ચ, ડિલીટ અને એક્સેલ એક્સપોર્ટ ફ્રેમ
frame_action = tk.Frame(root, padx=20)
frame_action.pack(fill=tk.X, pady=5)

tk.Label(frame_action, text="કસ્ટમર શોધો (નામ/નંબર/ઘર નં.):", font=("Arial", 11, "bold")).pack(side=tk.LEFT, padx=5)
entry_search = tk.Entry(frame_action, font=("Arial", 11), width=25)
entry_search.pack(side=tk.LEFT, padx=5)

btn_search = tk.Button(frame_action, text="શોધો", command=search_customer, bg="#17a2b8", fg="white", font=("Arial", 10, "bold"), width=10)
btn_search.pack(side=tk.LEFT, padx=5)

btn_reset = tk.Button(frame_action, text="બધા બતાવો", command=lambda: (entry_search.delete(0, tk.END), show_customers()), font=("Arial", 10), width=10)
btn_reset.pack(side=tk.LEFT, padx=5)

btn_export = tk.Button(frame_action, text="Export to Excel", command=export_to_excel, bg="#107c41", fg="white", font=("Arial", 10, "bold"))
btn_export.pack(side=tk.RIGHT, padx=5)

btn_delete = tk.Button(frame_action, text="પસંદ કરેલ ડિલીટ કરો", command=delete_customer, bg="#dc3545", fg="white", font=("Arial", 10, "bold"))
btn_delete.pack(side=tk.RIGHT, padx=5)

# ટેબલ ફ્રેમ
frame_table = tk.Frame(root, padx=20, pady=10)
frame_table.pack(fill=tk.BOTH, expand=True)

style = ttk.Style()
style.configure("Treeview.Heading", font=("Arial", 11, "bold"))
style.configure("Treeview", font=("Arial", 10), rowheight=28)

tree = ttk.Treeview(frame_table, columns=("No", "Name", "Phone", "HomeNo", "Amount"), show="headings")
tree.heading("No", text="ક્રમ નં.")
tree.heading("Name", text="કસ્ટમરનું નામ")
tree.heading("Phone", text="મોબાઇલ નંબર")
tree.heading("HomeNo", text="ઘર નંબર (Home No)")
tree.heading("Amount", text="રકમ (RS)")

tree.column("No", width=60, anchor="center")
tree.column("Name", width=300, anchor="w")
tree.column("Phone", width=180, anchor="center")
tree.column("HomeNo", width=150, anchor="center")
tree.column("Amount", width=150, anchor="center")

tree.pack(fill=tk.BOTH, expand=True)

# કુલ રકમ નો લેબલ
frame_total = tk.Frame(root, padx=20, pady=10)
frame_total.pack(fill=tk.X)

label_total = tk.Label(frame_total, text="કુલ રકમ: ₹ 0.00", font=("Arial", 13, "bold"), fg="#28a745")
label_total.pack(side=tk.RIGHT, padx=10)

# ડેટા લોડ કરો
show_customers()

root.mainloop()
